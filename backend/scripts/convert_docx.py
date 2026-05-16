"""Convert .docx files in backend/assets/ into base_resume.md / base_cover.md.

This is resume-aware: it walks the .docx via python-docx, drops embedded images,
and promotes bold/large paragraphs to markdown headings so the AI sees real
structure (## Experience, ## Skills, etc.) instead of an undifferentiated blob.

Usage (from backend/, venv active):
  python -m scripts.convert_docx                       # auto: resume.docx + cover.docx
  python -m scripts.convert_docx my_resume.docx cover.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

BACKEND_DIR = Path(__file__).resolve().parent.parent
ASSETS_DIR = BACKEND_DIR / "assets"

NAME_MIN_SIZE_PT = 14  # any bold run >= this becomes the H1 (the name)
SECTION_MAX_CHARS = 80  # bold paragraphs shorter than this become H2 sections


def _wrap(text: str, marker: str) -> str:
    """Wrap `text` with `marker` but keep any leading/trailing whitespace OUTSIDE it.
    CommonMark forbids whitespace adjacent to a strong/emphasis closing marker, so
    "**foo: **bar" parses as literal asterisks. We emit "**foo:** bar" instead."""
    leading_len = len(text) - len(text.lstrip())
    trailing_len = len(text) - len(text.rstrip())
    leading = text[:leading_len]
    trailing = text[len(text) - trailing_len:] if trailing_len else ""
    core = text.strip()
    if not core:
        return text
    return f"{leading}{marker}{core}{marker}{trailing}"


def _inline_md(paragraph) -> str:
    """Render a paragraph's runs as markdown with **bold** / *italic*, dropping images."""
    out: list[str] = []
    for r in paragraph.runs:
        text = r.text
        if not text:
            continue
        if r.bold:
            text = _wrap(text, "**")
        elif r.italic:
            text = _wrap(text, "*")
        out.append(text)
    return "".join(out).strip()


def _is_all_bold(paragraph) -> bool:
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        return False
    return all(r.bold for r in runs)


def _max_font_size(paragraph) -> float | None:
    sizes = [r.font.size.pt for r in paragraph.runs if r.font.size]
    return max(sizes) if sizes else None


def docx_to_markdown(path: Path) -> str:
    doc = Document(str(path))
    lines: list[str] = []
    seen_name = False
    contact_lines_after_name = 0  # absorb the 2-3 contact lines under the name

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            if lines and lines[-1] != "":
                lines.append("")
            continue

        # Strip out any paragraph that is just embedded image placeholders
        if text.startswith("![") and text.endswith(")"):
            continue

        size = _max_font_size(p)
        all_bold = _is_all_bold(p)
        is_list = (p.style and p.style.name == "List Paragraph")
        centered = p.alignment == WD_PARAGRAPH_ALIGNMENT.CENTER

        # H1 — name (large bold, usually first big thing)
        if not seen_name and all_bold and size and size >= NAME_MIN_SIZE_PT:
            lines.append(f"# {text}")
            lines.append("")
            seen_name = True
            contact_lines_after_name = 3
            continue

        # Right after the name: keep title + contact lines as plain paragraphs
        if contact_lines_after_name > 0 and centered:
            lines.append(text)
            contact_lines_after_name -= 1
            continue

        # H2 — section header (bold-only, short, NOT centered, NOT a list)
        if all_bold and not is_list and not centered and len(text) <= SECTION_MAX_CHARS:
            if lines and lines[-1] != "":
                lines.append("")
            lines.append(f"## {text}")
            lines.append("")
            continue

        # Bullet
        if is_list:
            lines.append(f"- {_inline_md(p)}")
            continue

        # Regular paragraph
        lines.append(_inline_md(p))

    return _tidy("\n".join(lines))


def _tidy(md: str) -> str:
    md = re.sub(r"\n{3,}", "\n\n", md).strip() + "\n"
    return md


def _resolve(arg: str | None, default_name: str) -> Path:
    if arg:
        p = Path(arg)
        if not p.is_absolute():
            p = ASSETS_DIR / p
        if not p.exists():
            sys.exit(f"Not found: {p}")
        return p

    for name in (default_name, default_name.replace("_", " ")):
        for ext in (".docx", ".DOCX"):
            candidate = ASSETS_DIR / f"{name}{ext}"
            if candidate.exists():
                return candidate
    sys.exit(
        f"No file given and no default found. Place '{default_name}.docx' in {ASSETS_DIR} "
        "or pass a path."
    )


def main() -> None:
    args = sys.argv[1:]
    resume_src = _resolve(args[0] if len(args) > 0 else None, "resume")
    cover_src = _resolve(args[1] if len(args) > 1 else None, "cover")

    print(f"Reading resume: {resume_src}")
    resume_md = docx_to_markdown(resume_src)
    (ASSETS_DIR / "base_resume.md").write_text(resume_md, encoding="utf-8")
    print(f"  -> wrote base_resume.md ({len(resume_md)} chars)")

    print(f"Reading cover:  {cover_src}")
    cover_md = docx_to_markdown(cover_src)
    (ASSETS_DIR / "base_cover.md").write_text(cover_md, encoding="utf-8")
    print(f"  -> wrote base_cover.md ({len(cover_md)} chars)")

    print("\nDone. Skim the .md files to confirm headings look right.")


if __name__ == "__main__":
    main()
